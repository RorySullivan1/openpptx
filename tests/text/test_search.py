# pyright: reportPrivateUsage=false

"""Unit-test suite for find/replace across runs: `pptx.text.search` and its `TextFrame` and
`_Paragraph` methods."""

from __future__ import annotations

import re
from typing import cast

import pytest

from pptx.exc import InvalidTypeError, InvalidValueError
from pptx.oxml.text import CT_TextBody, CT_TextParagraph
from pptx.text.search import TextChar, TextMatch
from pptx.text.text import TextFrame, _Paragraph

from ..unitutil.cxml import element, xml

# -- "Hello world!" split over three runs, each with its own formatting --
HELLO_P = 'a:p/(a:r/(a:rPr{b=1},a:t"Hel"),a:r/(a:rPr{i=1},a:t"lo wo"),a:r/(a:rPr{u=sng},a:t"rld!"))'


def _paragraph(p_cxml: str) -> _Paragraph:
    return _Paragraph(cast(CT_TextParagraph, element(p_cxml)), None)  # pyright: ignore


def _text_frame(txBody_cxml: str) -> TextFrame:
    return TextFrame(cast(CT_TextBody, element(txBody_cxml)), None)  # pyright: ignore


class Describe_ParagraphReplace:
    """Unit-test suite for `_Paragraph.replace()`."""

    @pytest.mark.parametrize(
        ("p_cxml", "old", "new", "expected_cxml"),
        [
            # -- within one run: splice its text, formatting untouched --
            (
                'a:p/a:r/(a:rPr{b=1},a:t"Hello world")',
                "world",
                "there",
                'a:p/a:r/(a:rPr{b=1},a:t"Hello there")',
            ),
            (
                'a:p/(a:r/(a:rPr{b=1},a:t"ab"),a:r/(a:rPr{i=1},a:t"cXd"))',
                "X",
                "YY",
                'a:p/(a:r/(a:rPr{b=1},a:t"ab"),a:r/(a:rPr{i=1},a:t"cYYd"))',
            ),
            # -- across runs: prefix + new in the first run, middle removed, suffix kept --
            (
                HELLO_P,
                "llo wor",
                "NEW",
                'a:p/(a:r/(a:rPr{b=1},a:t"HeNEW"),a:r/(a:rPr{u=sng},a:t"ld!"))',
            ),
            # -- a match starting at a run boundary takes that run's formatting --
            (
                HELLO_P,
                "lo world",
                "XY",
                'a:p/(a:r/(a:rPr{b=1},a:t"Hel"),a:r/(a:rPr{i=1},a:t"XY"),'
                'a:r/(a:rPr{u=sng},a:t"!"))',
            ),
            # -- a run left empty is removed --
            (
                HELLO_P,
                "rld!",
                "",
                'a:p/(a:r/(a:rPr{b=1},a:t"Hel"),a:r/(a:rPr{i=1},a:t"lo wo"))',
            ),
            (
                HELLO_P,
                "Hello world!",
                "",
                "a:p",
            ),
            (
                'a:p/(a:r/a:t"foo",a:endParaRPr{sz=1800})',
                "foo",
                "",
                "a:p/a:endParaRPr{sz=1800}",
            ),
            # -- an empty run inside the match goes with the middle runs --
            (
                'a:p/(a:r/a:t"ab",a:r/(a:rPr{b=1},a:t),a:r/a:t"cd")',
                "bc",
                "X",
                'a:p/(a:r/a:t"aX",a:r/a:t"d")',
            ),
            # -- every match in a run, and matches that share a run --
            (
                'a:p/(a:r/a:t"a-a",a:r/(a:rPr{b=1},a:t"-a-"))',
                "-",
                "+",
                'a:p/(a:r/a:t"a+a",a:r/(a:rPr{b=1},a:t"+a+"))',
            ),
            (
                'a:p/(a:r/a:t"xab",a:r/(a:rPr{b=1},a:t"cab"),a:r/(a:rPr{i=1},a:t"cz"))',
                "abc",
                "_",
                'a:p/(a:r/a:t"x_",a:r/(a:rPr{b=1},a:t"_"),a:r/(a:rPr{i=1},a:t"z"))',
            ),
        ],
    )
    def it_replaces_text_across_runs(self, p_cxml: str, old: str, new: str, expected_cxml: str):
        paragraph = _paragraph(p_cxml)
        old_text = paragraph.text

        count = paragraph.replace(old, new)

        assert paragraph._p.xml == xml(expected_cxml)
        assert paragraph.text == old_text.replace(old, new)
        assert count == old_text.count(old)

    def it_returns_zero_and_leaves_the_xml_alone_when_nothing_matches(self):
        paragraph = _paragraph(HELLO_P)
        before = paragraph._p.xml

        assert paragraph.replace("xyz", "abc") == 0
        assert paragraph._p.xml == before

    @pytest.mark.parametrize(
        "p_cxml",
        [
            'a:p/(a:r/a:t"foo",a:br,a:r/a:t"bar")',
            'a:p/(a:r/a:t"foo",a:fld{type=slidenum}/a:t"3",a:r/a:t"bar")',
        ],
    )
    @pytest.mark.parametrize("old", ["o\vb", "o3b", "ob", "foo\vbar", "3"])
    def it_does_not_match_across_a_line_break_or_a_field(self, p_cxml: str, old: str):
        paragraph = _paragraph(p_cxml)
        before = paragraph._p.xml

        assert paragraph.replace(old, "X") == 0
        assert paragraph._p.xml == before

    def but_it_matches_on_either_side_of_one(self):
        paragraph = _paragraph('a:p/(a:r/a:t"foo",a:fld{type=slidenum}/a:t"oo",a:br,a:r/a:t"foo")')

        assert paragraph.replace("oo", "0") == 2
        assert paragraph.text == "f0oo\vf0"

    def it_treats_the_pattern_literally_by_default(self):
        paragraph = _paragraph('a:p/a:r/a:t"a.b axb"')

        assert paragraph.replace(".", "-") == 1
        assert paragraph.text == "a-b axb"

    def it_treats_the_pattern_as_a_regex_when_asked(self):
        paragraph = _paragraph('a:p/(a:r/a:t"a1b2",a:r/(a:rPr{b=1},a:t"3c"))')

        assert paragraph.replace(r"\d+", "#", regex=True) == 2
        assert paragraph._p.xml == xml('a:p/(a:r/a:t"a#b#",a:r/(a:rPr{b=1},a:t"c"))')

    def it_expands_the_replacement_as_a_template_in_regex_mode(self):
        paragraph = _paragraph('a:p/(a:r/a:t"mail bob@exa",a:r/(a:rPr{b=1},a:t"mple.com now"))')

        count = paragraph.replace(r"(?P<user>\w+)@(\w+)\.com", r"[\g<user> at \2]", regex=True)

        assert count == 1
        assert paragraph._p.xml == xml(
            'a:p/(a:r/a:t"mail [bob at example]",a:r/(a:rPr{b=1},a:t" now"))'
        )

    def it_accepts_a_compiled_pattern_and_keeps_its_flags(self):
        paragraph = _paragraph('a:p/a:r/a:t"Foo foo FOO"')

        count = paragraph.replace(re.compile(r"(f)oo", re.IGNORECASE), r"\1--")

        assert count == 3
        assert paragraph.text == "F-- f-- F--"

    def but_it_does_not_expand_a_template_for_a_literal_pattern(self):
        paragraph = _paragraph('a:p/a:r/a:t"foo"')

        paragraph.replace("foo", r"\1 \g<x>")

        assert paragraph.text == r"\1 \g<x>"

    def it_skips_zero_width_regex_matches(self):
        paragraph = _paragraph('a:p/(a:r/a:t"ba",a:r/(a:rPr{b=1},a:t"ac"))')

        assert paragraph.replace("a*", "X", regex=True) == 1
        assert paragraph._p.xml == xml('a:p/(a:r/a:t"bX",a:r/(a:rPr{b=1},a:t"c"))')
        assert paragraph.replace(r"\b", "X", regex=True) == 0
        assert paragraph.replace("", "X", regex=True) == 0

    @pytest.mark.parametrize(
        ("count", "expected_count", "expected_text"),
        [(None, 3, "x-x-x"), (1, 1, "x-a-a"), (2, 2, "x-x-a"), (5, 3, "x-x-x"), (0, 0, "a-a-a")],
    )
    def it_can_limit_the_number_of_replacements(
        self, count: int | None, expected_count: int, expected_text: str
    ):
        paragraph = _paragraph('a:p/(a:r/a:t"a-",a:r/(a:rPr{b=1},a:t"a-a"))')

        assert paragraph.replace("a", "x", count=count) == expected_count
        assert paragraph.text == expected_text

    def it_raises_on_a_negative_count(self):
        paragraph = _paragraph('a:p/a:r/a:t"a"')
        with pytest.raises(InvalidValueError, match="count"):
            paragraph.replace("a", "b", count=-1)

    def it_raises_on_an_empty_literal_pattern(self):
        paragraph = _paragraph('a:p/a:r/a:t"a"')
        with pytest.raises(InvalidValueError, match="empty"):
            paragraph.replace("", "b")

    def it_raises_on_an_invalid_template_before_changing_anything(self):
        paragraph = _paragraph('a:p/(a:r/a:t"ab",a:r/(a:rPr{b=1},a:t"ab"))')
        before = paragraph._p.xml
        with pytest.raises(InvalidValueError, match="template"):
            paragraph.replace("(a)", r"\9", regex=True)
        assert paragraph._p.xml == before

    @pytest.mark.parametrize("new", [r"\2", r"\g<foo>", "\\"])
    def it_raises_on_an_invalid_template_even_when_nothing_matches(self, new: str):
        paragraph = _paragraph('a:p/a:r/a:t"no digits here"')
        with pytest.raises(InvalidValueError, match="template"):
            paragraph.replace(r"(\d+)", new, regex=True)

    def and_it_raises_on_a_template_naming_an_unknown_group_where_one_matches(self):
        paragraph = _paragraph('a:p/a:r/a:t"abc"')
        with pytest.raises(InvalidValueError, match="template"):
            paragraph.replace("(a)", r"\g<foo>", regex=True)
        assert paragraph.text == "abc"

    @pytest.mark.parametrize("count", ["2", 1.5, True])
    def it_raises_on_a_count_that_is_not_an_int(self, count: object):
        paragraph = _paragraph('a:p/a:r/a:t"aaa"')
        with pytest.raises(InvalidTypeError, match="count"):
            paragraph.replace("a", "b", count=count)  # pyright: ignore[reportArgumentType]
        assert paragraph.text == "aaa"

    @pytest.mark.parametrize(
        ("p_cxml", "old", "new", "expected_cxml"),
        [
            # -- within one run: the text after the match follows the last line --
            (
                'a:p/a:r/(a:rPr{b=1},a:t"aXb")',
                "X",
                "1\v2",
                'a:p/(a:r/(a:rPr{b=1},a:t"a1"),a:br/a:rPr{b=1},a:r/(a:rPr{b=1},a:t"2b"))',
            ),
            # -- a line feed is a line break too, as in the `_Paragraph.text` setter --
            (
                'a:p/a:r/(a:rPr{b=1},a:t"aXb")',
                "X",
                "1\n2",
                'a:p/(a:r/(a:rPr{b=1},a:t"a1"),a:br/a:rPr{b=1},a:r/(a:rPr{b=1},a:t"2b"))',
            ),
            # -- across runs: the new lines take the first run's formatting, the suffix keeps its own --
            (
                HELLO_P,
                "llo wor",
                "A\vB",
                'a:p/(a:r/(a:rPr{b=1},a:t"HeA"),a:br/a:rPr{b=1},a:r/(a:rPr{b=1},a:t"B"),'
                'a:r/(a:rPr{u=sng},a:t"ld!"))',
            ),
            # -- an empty line gets its break but no run, and an emptied run is removed --
            (
                'a:p/a:r/(a:rPr{b=1},a:t"X")',
                "X",
                "\v",
                "a:p/a:br/a:rPr{b=1}",
            ),
            (
                'a:p/a:r/a:t"aXb"',
                "X",
                "\v\v",
                'a:p/(a:r/a:t"a",a:br,a:br,a:r/a:t"b")',
            ),
        ],
    )
    def it_inserts_a_line_break_for_each_line_break_in_the_new_text(
        self, p_cxml: str, old: str, new: str, expected_cxml: str
    ):
        paragraph = _paragraph(p_cxml)
        expected_text = paragraph.text.replace(old, new.replace("\n", "\v"))

        assert paragraph.replace(old, new) == 1

        assert paragraph._p.xml == xml(expected_cxml)
        assert paragraph.text == expected_text

    def and_it_inserts_line_breaks_from_an_expanded_template(self):
        paragraph = _paragraph('a:p/(a:r/a:t"one two",a:r/a:t" three")')

        assert paragraph.replace(r"(\w+) (\w+)", r"\2\v\1", regex=True) == 1

        assert paragraph.text == "two\vone three"

    @pytest.mark.parametrize(
        ("pattern", "expected"),
        [
            # -- `^` matches only at the start of the paragraph, never after a field --
            (r"^end", []),
            (r"^page", [(0, 4)]),
            # -- `\b` and lookbehind at a segment start see the text before it --
            (r"\bend", []),
            (r"(?<=3)end", [(6, 9)]),
            (r"(?<!3)end", []),
        ],
    )
    def it_matches_a_regex_in_the_context_of_the_paragraph_text(
        self, pattern: str, expected: list[tuple[int, int]]
    ):
        paragraph = _paragraph('a:p/(a:r/a:t"page ",a:fld{type=slidenum}/a:t"3",a:r/a:t"end")')
        assert paragraph.text == "page 3end"

        matches = paragraph.find_all(pattern, regex=True)

        assert [(m.start, m.end) for m in matches] == expected
        assert paragraph.replace(pattern, "X", regex=True) == len(expected)

    def and_it_treats_the_end_of_a_segment_as_the_end_of_the_text(self):
        paragraph = _paragraph('a:p/(a:r/a:t"one",a:br,a:r/a:t"two")')

        assert [m.text for m in paragraph.find_all(r"\w+$", regex=True)] == ["one", "two"]
        assert [m.text for m in paragraph.find_all(r"^\w+", regex=True)] == ["one"]

    def it_raises_on_an_invalid_regex(self):
        paragraph = _paragraph('a:p/a:r/a:t"a"')
        with pytest.raises(InvalidValueError, match="regular expression"):
            paragraph.replace("(", "b", regex=True)

    @pytest.mark.parametrize(("old", "new"), [(None, "b"), (b"a", "b"), ("a", None), ("a", 1)])
    def it_raises_on_a_pattern_or_replacement_of_the_wrong_type(self, old: object, new: object):
        paragraph = _paragraph('a:p/a:r/a:t"a"')
        with pytest.raises(InvalidTypeError):
            paragraph.replace(old, new)  # pyright: ignore[reportArgumentType]


class Describe_ParagraphFindAll:
    """Unit-test suite for `_Paragraph.find_all()`."""

    def it_finds_each_match_with_its_offsets_and_runs(self):
        paragraph = _paragraph(HELLO_P)
        r_lst = paragraph._p.r_lst

        matches = paragraph.find_all("o")

        assert [(m.start, m.end, m.text) for m in matches] == [(4, 5, "o"), (7, 8, "o")]
        assert [[r._r for r in m.runs] for m in matches] == [[r_lst[1]], [r_lst[1]]]

    def it_lists_exactly_the_runs_a_match_touches(self):
        paragraph = _paragraph('a:p/(a:r/a:t"ab",a:r/a:t,a:r/a:t"cd",a:r/a:t"ef")')
        r_lst = paragraph._p.r_lst

        (match,) = paragraph.find_all("bc")

        assert [r._r for r in match.runs] == [r_lst[0], r_lst[2]]

    @pytest.mark.parametrize(
        ("p_cxml", "pattern", "regex"),
        [
            (HELLO_P, "llo wor", False),
            (HELLO_P, r"\w+", True),
            ('a:p/(a:fld/a:t"12",a:r/a:t"ab",a:br,a:r/a:t"abc",a:r/a:t"ab")', "b", False),
            ('a:p/(a:fld/a:t"12",a:r/a:t"ab",a:br,a:r/a:t"abc",a:r/a:t"ab")', "ca", False),
        ],
    )
    def it_offsets_each_match_into_the_paragraph_text(
        self, p_cxml: str, pattern: str, regex: bool
    ):
        paragraph = _paragraph(p_cxml)
        text = paragraph.text

        matches = paragraph.find_all(pattern, regex=regex)

        assert matches
        for m in matches:
            assert text[m.start : m.end] == m.text
            assert m.paragraph is paragraph
            assert m.paragraph_index == 0
            assert m.shape is None
            assert "".join(r.text for r in m.runs).find(m.text) != -1

    def it_finds_nothing_in_a_paragraph_without_runs(self):
        assert _paragraph('a:p/(a:fld/a:t"foo",a:br)').find_all("foo") == []

    def it_does_not_change_the_xml(self):
        paragraph = _paragraph(HELLO_P)
        before = paragraph._p.xml

        paragraph.find_all("l", regex=True)

        assert paragraph._p.xml == before

    def it_returns_immutable_matches(self):
        (match,) = _paragraph('a:p/a:r/a:t"foo"').find_all("foo")

        assert isinstance(match, TextMatch)
        with pytest.raises(AttributeError):
            match.start = 2  # pyright: ignore[reportAttributeAccessIssue]

    def it_compares_and_hashes_matches_by_what_they_refer_to(self):
        paragraph = _paragraph(HELLO_P)

        first, second = paragraph.find_all("o"), paragraph.find_all("o")

        assert first == second
        assert len({*first, *second}) == 2
        assert first[0] != first[1]
        # -- the same offsets in a different paragraph are a different match --
        assert _paragraph(HELLO_P).find_all("o") != first

    def it_raises_on_an_empty_literal_pattern(self):
        with pytest.raises(InvalidValueError):
            _paragraph('a:p/a:r/a:t"a"').find_all("")


class DescribeTextFrameFindReplace:
    """Unit-test suite for `TextFrame.find_all()` and `TextFrame.replace()`."""

    TX_BODY = (
        'p:txBody/(a:bodyPr,a:p/(a:r/a:t"one fi",a:r/(a:rPr{b=1},a:t"sh")),'
        'a:p/a:r/a:t"two fish",a:p/a:r/a:t"red fish")'
    )

    def it_finds_matches_in_every_paragraph(self):
        text_frame = _text_frame(self.TX_BODY)

        matches = text_frame.find_all("fish")

        assert [(m.paragraph_index, m.start) for m in matches] == [(0, 4), (1, 4), (2, 4)]
        paragraphs = text_frame.paragraphs
        for m in matches:
            assert m.paragraph._p is paragraphs[m.paragraph_index]._p

    def it_does_not_match_across_paragraphs(self):
        text_frame = _text_frame(self.TX_BODY)
        assert text_frame.find_all("sh\ntwo") == []
        assert text_frame.find_all("shtwo") == []

    @pytest.mark.parametrize(
        ("count", "expected_count", "expected_text"),
        [
            (None, 3, "one cat\ntwo cat\nred cat"),
            (2, 2, "one cat\ntwo cat\nred fish"),
            (0, 0, "one fish\ntwo fish\nred fish"),
        ],
    )
    def it_replaces_across_the_frame_in_document_order(
        self, count: int | None, expected_count: int, expected_text: str
    ):
        text_frame = _text_frame(self.TX_BODY)

        assert text_frame.replace("fish", "cat", count=count) == expected_count
        assert text_frame.text == expected_text

    def it_gives_the_new_text_the_first_runs_formatting(self):
        text_frame = _text_frame(self.TX_BODY)

        text_frame.replace("fish", "cat", count=1)

        runs = text_frame.paragraphs[0].runs
        assert [(r.text, r._r.formatting_key()) for r in runs] == [("one cat", ())]


class DescribeTextFrameTraversal:
    """Unit-test suite for `iter_runs()` and `iter_chars()`."""

    def it_iterates_every_run_in_document_order(self):
        text_frame = _text_frame(
            'p:txBody/(a:bodyPr,a:p/(a:r/a:t"a",a:fld/a:t"1",a:br,a:r/a:t"b"),a:p,a:p/a:r/a:t"c")'
        )

        runs = list(text_frame.iter_runs())

        assert [r.text for r in runs] == ["a", "b", "c"]
        assert runs[2]._parent._p is text_frame.paragraphs[2]._p  # pyright: ignore

    def it_iterates_every_character_with_its_position_and_run(self):
        text_frame = _text_frame(
            'p:txBody/(a:bodyPr,a:p/(a:r/a:t"ab",a:br,a:fld/a:t"7"),a:p/a:r/(a:rPr{b=1},a:t"c"))'
        )
        r0 = text_frame.paragraphs[0]._p.r_lst[0]
        r1 = text_frame.paragraphs[1]._p.r_lst[0]

        chars = list(text_frame.iter_chars())

        assert all(isinstance(c, TextChar) for c in chars)
        assert [(c.char, c.paragraph_index, c.offset) for c in chars] == [
            ("a", 0, 0),
            ("b", 0, 1),
            ("\v", 0, 2),
            ("7", 0, 3),
            ("c", 1, 0),
        ]
        assert [c.run and c.run._r for c in chars] == [r0, r0, None, None, r1]

    def it_iterates_the_characters_of_one_paragraph(self):
        paragraph = _paragraph('a:p/(a:r/a:t"x",a:r/a:t"yz")')

        chars = list(paragraph.iter_chars())

        assert "".join(c.char for c in chars) == paragraph.text
        assert [c.offset for c in chars] == [0, 1, 2]
        assert {c.paragraph_index for c in chars} == {0}


class DescribeMergeRuns:
    """Unit-test suite for `merge_runs()`."""

    @pytest.mark.parametrize(
        ("p_cxml", "expected_count", "expected_cxml"),
        [
            ("a:p", 0, "a:p"),
            ('a:p/a:r/a:t"a"', 0, 'a:p/a:r/a:t"a"'),
            ('a:p/(a:r/a:t"a",a:r/a:t"b")', 1, 'a:p/a:r/a:t"ab"'),
            # -- proofing state is ignored, the first run's properties are kept --
            (
                'a:p/(a:r/(a:rPr{lang=en-US,dirty=0},a:t"a"),a:r/(a:rPr{lang=en-US,err=1},a:t"b"),'
                'a:r/(a:rPr{lang=en-US,smtClean=0,smtId=3},a:t"c"))',
                2,
                'a:p/a:r/(a:rPr{lang=en-US,dirty=0},a:t"abc")',
            ),
            ('a:p/(a:r/(a:rPr{dirty=0},a:t"a"),a:r/a:t"b")', 1, 'a:p/a:r/(a:rPr{dirty=0},a:t"ab")'),
            # -- real formatting differences are kept --
            (
                'a:p/(a:r/(a:rPr{b=1},a:t"a"),a:r/a:t"b",a:r/(a:rPr{i=1},a:t"c"),'
                'a:r/(a:rPr{sz=1200},a:t"d"),a:r/(a:rPr{baseline=30000},a:t"e"))',
                0,
                'a:p/(a:r/(a:rPr{b=1},a:t"a"),a:r/a:t"b",a:r/(a:rPr{i=1},a:t"c"),'
                'a:r/(a:rPr{sz=1200},a:t"d"),a:r/(a:rPr{baseline=30000},a:t"e"))',
            ),
            # -- runs either side of a line break or a field are not adjacent --
            (
                'a:p/(a:r/a:t"a",a:br,a:r/a:t"b",a:fld/a:t"1",a:r/a:t"c",a:r/a:t"d")',
                1,
                'a:p/(a:r/a:t"a",a:br,a:r/a:t"b",a:fld/a:t"1",a:r/a:t"cd")',
            ),
            (
                'a:p/(a:r/(a:rPr{b=1},a:t"a"),a:r/(a:rPr{b=1},a:t"b"),a:r/a:t"c",'
                'a:r/a:t"d",a:r/(a:rPr{b=1},a:t"e"))',
                2,
                'a:p/(a:r/(a:rPr{b=1},a:t"ab"),a:r/a:t"cd",a:r/(a:rPr{b=1},a:t"e"))',
            ),
        ],
    )
    def it_merges_adjacent_runs_with_the_same_formatting(
        self, p_cxml: str, expected_count: int, expected_cxml: str
    ):
        paragraph = _paragraph(p_cxml)
        text = paragraph.text

        assert paragraph.merge_runs() == expected_count
        assert paragraph._p.xml == xml(expected_cxml)
        assert paragraph.text == text

    def it_merges_runs_in_every_paragraph_of_a_frame(self):
        text_frame = _text_frame(
            'p:txBody/(a:bodyPr,a:p/(a:r/a:t"a",a:r/a:t"b"),a:p/(a:r/a:t"c",a:r/a:t"d"))'
        )

        assert text_frame.merge_runs() == 2
        assert [len(p.runs) for p in text_frame.paragraphs] == [1, 1]
