# pyright: reportPrivateUsage=false

"""Integration tests for deck-wide find/replace: `Slide` and `Presentation` `find_text()` and
`replace_text()`, run against real presentations, saved and reopened."""

from __future__ import annotations

import io
from pathlib import Path

import pytest

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.exc import InvalidValueError
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from pptx.presentation import Presentation as PresentationCls
from pptx.shapes.graphfrm import GraphicFrame
from pptx.slide import Slide
from pptx.util import Inches

CORPUS_DIR = Path(__file__).parent / "test_files" / "real_world"
URL = "https://example.com/widgets"


def _reopen(prs: PresentationCls) -> PresentationCls:
    stream = io.BytesIO()
    prs.save(stream)
    stream.seek(0)
    return Presentation(stream)


def _add_split_textbox(slide: Slide, *texts: str):
    """A text box whose one paragraph is `texts`, a run each, alternately bold."""
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    paragraph = textbox.text_frame.paragraphs[0]
    for idx, text in enumerate(texts):
        run = paragraph.add_run()
        run.text = text
        run.font.bold = idx % 2 == 0
    return textbox


@pytest.fixture
def prs() -> PresentationCls:
    """Two slides; the first has text in a text box, a group, a table and its notes."""
    prs = Presentation()
    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)
    textbox = _add_split_textbox(slide, "Acme Wid", "gets", " rule")
    textbox.name = "split"
    group = slide.shapes.add_group_shape()
    member = group.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    member.name = "member"
    member.text_frame.text = "Widgets in a group"
    graphic_frame = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(4), Inches(1))
    graphic_frame.name = "table"
    graphic_frame.table.cell(0, 1).text = "Widgets in a cell"
    graphic_frame.table.cell(1, 0).text = "More Widgets"
    slide.notes_slide.notes_text_frame.text = "Widgets in the notes"  # pyright: ignore

    second = prs.slides.add_slide(blank)
    second.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "Widgets, again"
    return prs


class DescribeSlideFindText:
    """`Slide.find_text()`."""

    def it_covers_group_members_table_cells_and_the_notes(self, prs: PresentationCls):
        slide = prs.slides[0]

        matches = slide.find_text("Widgets")

        assert [m.paragraph.text for m in matches] == [
            "Acme Widgets rule",
            "Widgets in a group",
            "Widgets in a cell",
            "More Widgets",
            "Widgets in the notes",
        ]
        assert [m.shape.name for m in matches[:4]] == ["split", "member", "table", "table"]
        assert isinstance(matches[2].shape, GraphicFrame)
        notes_shape = slide.notes_slide.notes_placeholder
        assert notes_shape is not None and matches[4].shape == notes_shape
        assert [r.text for r in matches[0].runs] == ["Acme Wid", "gets"]

    def it_does_not_create_a_notes_slide(self, prs: PresentationCls):
        second = prs.slides[1]

        second.find_text("Widgets")
        second.replace_text("Widgets", "Gadgets")

        assert second.has_notes_slide is False

    def it_does_not_add_a_text_body_to_a_shape_or_cell_without_one(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(1), Inches(1))
        shape._sp.remove(shape._sp.txBody)
        tbl = slide.shapes.add_table(1, 1, 0, 0, Inches(1), Inches(1)).table
        tc = tbl._tbl.tr_lst[0].tc_lst[0]
        tc.remove(tc.txBody)
        before = slide.part.blob

        assert slide.find_text("x") == []
        assert slide.replace_text("x", "y") == 0
        assert slide.part.blob == before

    def it_does_not_match_inherited_placeholder_prompt_text(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])

        assert slide.find_text("Click", regex=True) == []

    def it_raises_on_an_empty_pattern_even_with_no_text_to_search(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        with pytest.raises(InvalidValueError):
            slide.find_text("")
        with pytest.raises(InvalidValueError):
            prs.replace_text("", "x")


class DescribePresentationFindReplace:
    """`Presentation.find_text()` and `Presentation.replace_text()`."""

    def it_finds_across_slides_in_presentation_order(self, prs: PresentationCls):
        matches = prs.find_text("Widgets")

        assert len(matches) == 6
        assert matches[-1].paragraph.text == "Widgets, again"

    @pytest.mark.parametrize(
        ("count", "expected_count", "expected_left"),
        [(None, 6, 0), (6, 6, 0), (3, 3, 3), (5, 5, 1), (0, 0, 6)],
    )
    def it_limits_replacements_across_slides_in_order(
        self, prs: PresentationCls, count: int | None, expected_count: int, expected_left: int
    ):
        assert prs.replace_text("Widgets", "Gadgets", count=count) == expected_count

        left = prs.find_text("Widgets")
        assert len(left) == expected_left
        # -- what is left is the tail of the deck, in order --
        assert [m.paragraph.text for m in left] == [
            m.paragraph.text for m in prs.find_text("Widgets|Gadgets", regex=True)
        ][expected_count:] or not left

    def it_replaces_a_word_split_over_runs_with_the_first_runs_formatting(
        self, prs: PresentationCls
    ):
        prs.replace_text("Widgets", "Gadgets")

        runs = prs.slides[0].shapes[0].text_frame.paragraphs[0].runs
        assert [(r.text, r.font.bold) for r in runs] == [("Acme Gadgets", True), (" rule", True)]

    def it_leaves_layouts_and_masters_byte_identical(self, prs: PresentationCls):
        master = prs.slide_masters[0]
        parts = [master.part, *(layout.part for layout in master.slide_layouts)]
        before = [part.blob for part in parts]
        assert "Click to edit Master title style" in master.part.blob.decode("utf-8")

        assert prs.replace_text("Click", "Tap") == 0
        assert prs.replace_text("Widgets", "Gadgets") == 6

        assert [part.blob for part in parts] == before

    def it_survives_a_save_and_reload(self, prs: PresentationCls):
        prs.replace_text(r"W(idget)s", r"G\1-s", regex=True)

        reopened = _reopen(prs)

        assert reopened.find_text("Widgets") == []
        assert len(reopened.find_text("Gidget-s")) == 6
        assert reopened.replace_text("Widgets", "x") == 0


class DescribeHyperlinkedRunRemoval:
    """A run removed by a replace takes its hyperlink relationship with it, when unshared."""

    def _deck(self, share: bool):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        textbox = _add_split_textbox(slide, "visit ", "our site", " today")
        textbox.text_frame.paragraphs[0].runs[1].hyperlink.address = URL
        if share:
            other = slide.shapes.add_textbox(0, 0, Inches(1), Inches(1))
            run = other.text_frame.paragraphs[0].add_run()
            run.text = "also linked"
            run.hyperlink.address = URL
        return prs, slide

    @staticmethod
    def _hyperlink_rIds(slide: Slide) -> list[str]:
        return [rId for rId, rel in slide.part.rels.items() if rel.reltype == RT.HYPERLINK]

    def it_drops_the_relationship_when_nothing_else_refers_to_it(self):
        prs, slide = self._deck(share=False)
        assert len(self._hyperlink_rIds(slide)) == 1

        assert slide.replace_text("visit our site", "see") == 1

        assert self._hyperlink_rIds(slide) == []
        assert slide.shapes[0].text_frame.text == "see today"
        assert _reopen(prs).slides[0].shapes[0].text_frame.text == "see today"

    def but_it_keeps_the_relationship_while_another_run_refers_to_it(self):
        prs, slide = self._deck(share=True)
        (rId,) = self._hyperlink_rIds(slide)

        assert slide.replace_text("visit our site", "see") == 1

        assert self._hyperlink_rIds(slide) == [rId]
        reopened = _reopen(prs).slides[0]
        assert reopened.shapes[1].text_frame.paragraphs[0].runs[0].hyperlink.address == URL

    def and_it_keeps_the_hyperlink_of_a_run_that_stays(self):
        prs, slide = self._deck(share=False)

        assert slide.replace_text("our", "my") == 1

        runs = slide.shapes[0].text_frame.paragraphs[0].runs
        assert [r.text for r in runs] == ["visit ", "my site", " today"]
        assert runs[1].hyperlink.address == URL

    def it_drops_a_relationship_the_removed_run_refers_to_twice(self):
        prs, slide = self._deck(share=False)
        hlinkClick = slide.shapes[0].text_frame.paragraphs[0].runs[1]._r.rPr.hlinkClick
        rId = hlinkClick.rId
        hlinkClick.addnext(parse_xml(f'<a:hlinkMouseOver {nsdecls("a", "r")} r:id="{rId}"/>'))

        assert slide.replace_text("visit our site", "see") == 1

        assert self._hyperlink_rIds(slide) == []

    def but_it_keeps_the_relationship_when_new_lines_carry_the_hyperlink_on(self):
        prs, slide = self._deck(share=False)
        (rId,) = self._hyperlink_rIds(slide)

        # -- the linked run empties and goes; the line after the break copies its formatting --
        assert slide.replace_text("our site", "\vour new site") == 1

        assert self._hyperlink_rIds(slide) == [rId]
        paragraph = _reopen(prs).slides[0].shapes[0].text_frame.paragraphs[0]
        assert paragraph.text == "visit \vour new site today"
        assert [r.hyperlink.address for r in paragraph.runs] == [None, URL, None]


class DescribeMatchesAndLineBreaks:
    """`TextMatch` values from a slide, and line breaks written by `replace_text()`."""

    def it_returns_matches_that_compare_and_hash_by_what_they_refer_to(
        self, prs: PresentationCls
    ):
        slide = prs.slides[0]

        first, second = slide.find_text("Widgets"), slide.find_text("Widgets")

        assert first == second
        assert len({*first, *second}) == 5

    def it_writes_a_line_break_for_each_one_in_the_new_text(self, prs: PresentationCls):
        assert prs.replace_text("Widgets", "Wid-\vgets") == 6

        reopened = _reopen(prs)

        texts = [m.paragraph.text for m in reopened.find_text("gets")]
        assert texts[0] == "Acme Wid-\vgets rule"
        assert all("\v" in text for text in texts)
        assert "_x000B_" not in reopened.slides[0].part.blob.decode("utf-8")


@pytest.mark.skipif(not CORPUS_DIR.is_dir(), reason="real-world corpus not present")
class DescribeCorpusRoundTrip:
    """A replace across runs in a PowerPoint-authored deck survives a save and reload."""

    def it_replaces_a_word_split_over_runs_in_a_corpus_deck(self):
        prs = Presentation(str(CORPUS_DIR / "customGeo.pptx"))
        before = prs.find_text("Assessments")
        assert any(len(m.runs) > 1 for m in before)

        count = prs.replace_text("Assessments", "Reviews")

        assert count == len(before)
        reopened = _reopen(prs)
        assert reopened.find_text("Assessments") == []
        assert len(reopened.find_text("Reviews")) >= count
        assert reopened.replace_text("Assessments", "Reviews") == 0
