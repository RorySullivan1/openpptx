"""Unit-test suite for `pptx.dml.effect` module."""

from __future__ import annotations

import io

import pytest

from pptx import Presentation
from pptx.dml.color import ColorFormat, RGBColor
from pptx.dml.effect import GlowFormat, ReflectionFormat, ShadowFormat, SoftEdgeFormat
from pptx.enum.dml import MSO_RECT_ALIGNMENT, MSO_SHADOW_TYPE, MSO_THEME_COLOR
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.util import Emu, Pt

from ..unitutil.cxml import element, xml


class DescribeShadowFormat:

    def it_adds_a_shadow_again_when_its_color_is_set_after_inheritance_is_restored(self):
        spPr = element("p:spPr")
        shadow = ShadowFormat(spPr)
        shadow.color.rgb = RGBColor(0xFF, 0x00, 0x00)
        shadow.inherit = True

        shadow.color.rgb = RGBColor(0x00, 0x00, 0xFF)

        assert spPr.xpath("a:effectLst/a:outerShdw/a:srgbClr/@val") == ["0000FF"]

    @pytest.mark.parametrize(
        ("spPr_cxml", "expected_value"),
        [
            ("p:spPr", None),
            ("p:spPr/a:effectLst", None),
            ("p:spPr/a:effectLst/a:outerShdw{blurRad=50800}", Emu(50800)),
        ],
    )
    def it_knows_its_blur_radius(self, spPr_cxml: str, expected_value):
        shadow = ShadowFormat(element(spPr_cxml))
        assert shadow.blur_radius == expected_value

    def it_can_change_its_blur_radius(self):
        spPr = element("p:spPr{a:b=c}")
        shadow = ShadowFormat(spPr)
        shadow.blur_radius = 50800
        assert spPr.xml == xml("p:spPr{a:b=c}/a:effectLst/a:outerShdw{blurRad=50800}/a:prstClr{val=black}/a:alpha{val=40000}")

    @pytest.mark.parametrize(
        ("spPr_cxml", "expected_value"),
        [
            ("p:spPr", None),
            ("p:spPr/a:effectLst/a:outerShdw{dist=38100}", Emu(38100)),
        ],
    )
    def it_knows_its_distance(self, spPr_cxml: str, expected_value):
        shadow = ShadowFormat(element(spPr_cxml))
        assert shadow.distance == expected_value

    def it_can_change_its_distance(self):
        spPr = element("p:spPr{a:b=c}")
        shadow = ShadowFormat(spPr)
        shadow.distance = 38100
        assert spPr.xml == xml("p:spPr{a:b=c}/a:effectLst/a:outerShdw{dist=38100}/a:prstClr{val=black}/a:alpha{val=40000}")

    @pytest.mark.parametrize(
        ("spPr_cxml", "expected_value"),
        [
            ("p:spPr", None),
            ("p:spPr/a:effectLst/a:outerShdw{dir=2700000}", 45.0),
        ],
    )
    def it_knows_its_direction(self, spPr_cxml: str, expected_value):
        shadow = ShadowFormat(element(spPr_cxml))
        assert shadow.direction == expected_value

    def it_can_change_its_direction(self):
        spPr = element("p:spPr{a:b=c}")
        shadow = ShadowFormat(spPr)
        shadow.direction = 45.0
        assert spPr.xml == xml("p:spPr{a:b=c}/a:effectLst/a:outerShdw{dir=2700000}/a:prstClr{val=black}/a:alpha{val=40000}")

    @pytest.mark.parametrize(
        ("spPr_cxml", "expected_value"),
        [
            ("p:spPr", None),
            ("p:spPr/a:effectLst/a:outerShdw{algn=bl}", MSO_RECT_ALIGNMENT.BOTTOM_LEFT),
        ],
    )
    def it_knows_its_alignment(self, spPr_cxml: str, expected_value):
        shadow = ShadowFormat(element(spPr_cxml))
        assert shadow.alignment == expected_value

    @pytest.mark.parametrize(
        ("spPr_cxml", "expected_value"),
        [
            ("p:spPr", None),
            ("p:spPr/a:effectLst/a:outerShdw", MSO_SHADOW_TYPE.OUTER),
            ("p:spPr/a:effectLst/a:innerShdw", MSO_SHADOW_TYPE.INNER),
        ],
    )
    def it_knows_its_shadow_type(
        self, spPr_cxml: str, expected_value: MSO_SHADOW_TYPE | None
    ):
        shadow = ShadowFormat(element(spPr_cxml))
        assert shadow.shadow_type == expected_value

    def it_knows_whether_it_inherits(self, inherit_get_fixture):
        shadow, expected_value = inherit_get_fixture
        inherit = shadow.inherit
        assert inherit is expected_value

    def it_can_change_whether_it_inherits(self, inherit_set_fixture):
        shadow, value, expected_xml = inherit_set_fixture
        shadow.inherit = value
        assert shadow._element.xml == expected_xml

    @pytest.mark.parametrize(
        ("prop_name", "value", "attr"),
        [
            ("blur_radius", Emu(50800), "blurRad=50800"),
            ("distance", Emu(38100), "dist=38100"),
            ("direction", 45.0, "dir=2700000"),
            ("alignment", MSO_RECT_ALIGNMENT.BOTTOM_LEFT, "algn=bl"),
            ("rotate_with_shape", False, "rotWithShape=0"),
        ],
    )
    def it_adds_a_default_color_when_it_creates_a_shadow(self, prop_name, value, attr):
        # -- `CT_OuterShadowEffect` requires a color child; PowerPoint refuses a deck without one (#131)
        spPr = element("p:spPr{a:b=c}")
        setattr(ShadowFormat(spPr), prop_name, value)
        assert spPr.xml == xml(
            "p:spPr{a:b=c}/a:effectLst/a:outerShdw{%s}/a:prstClr{val=black}/a:alpha{val=40000}" % attr
        )

    def it_keeps_an_existing_color_when_it_changes_a_shadow(self):
        spPr = element("p:spPr/a:effectLst/a:outerShdw{blurRad=50800}/a:srgbClr{val=FF0000}")
        ShadowFormat(spPr).blur_radius = Emu(101600)
        assert spPr.xml == xml(
            "p:spPr/a:effectLst/a:outerShdw{blurRad=101600}/a:srgbClr{val=FF0000}"
        )

    def it_repairs_a_colorless_shadow_when_it_changes_it(self):
        spPr = element("p:spPr/a:effectLst/a:outerShdw{blurRad=50800}")
        ShadowFormat(spPr).distance = Emu(38100)
        assert spPr.xml == xml(
            "p:spPr/a:effectLst/a:outerShdw{blurRad=50800,dist=38100}"
            "/a:prstClr{val=black}/a:alpha{val=40000}"
        )

    def it_round_trips_a_shadow_through_a_saved_deck(self):
        # -- regression for #131: the deck from the issue's repro must carry a shadow color --
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE, Emu(914400), Emu(914400), Emu(2743200), Emu(1828800)
        )
        box.shadow.blur_radius = Pt(8)
        other = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.OVAL, Emu(914400), Emu(3657600), Emu(2743200), Emu(1371600)
        )
        other.shadow.color.rgb = RGBColor(0x00, 0x80, 0xFF)
        other.shadow.distance = Pt(4)
        stream = io.BytesIO()
        prs.save(stream)

        stream.seek(0)
        shapes = Presentation(stream).slides[0].shapes
        reloaded_box, reloaded_other = shapes[-2], shapes[-1]
        assert reloaded_box.shadow.blur_radius == Pt(8)
        assert reloaded_box.shadow.color.alpha == pytest.approx(0.4)
        assert reloaded_box.shadow._outerShdw.xpath("a:prstClr/@val") == ["black"]
        assert reloaded_other.shadow.distance == Pt(4)
        assert reloaded_other.shadow.color.rgb == RGBColor(0x00, 0x80, 0xFF)

    def it_provides_access_to_shadow_color(self):
        from pptx.oxml import parse_xml
        from pptx.oxml.ns import nsdecls

        # Shadow with an sRGB color
        spPr_xml = (
            '<p:spPr %s><a:effectLst>'
            '<a:outerShdw blurRad="50800" dist="38100" dir="2700000">'
            '<a:srgbClr val="FF0000"/>'
            '</a:outerShdw></a:effectLst></p:spPr>' % nsdecls("p", "a")
        )
        spPr = parse_xml(spPr_xml)
        shadow = ShadowFormat(spPr)
        assert isinstance(shadow.color, ColorFormat)
        assert shadow.color.rgb == RGBColor(0xFF, 0x00, 0x00)

    def it_can_set_shadow_color(self):
        shadow = ShadowFormat(element("p:spPr"))
        shadow.color.rgb = RGBColor(0x00, 0x80, 0xFF)
        assert shadow.color.rgb == RGBColor(0x00, 0x80, 0xFF)
        assert shadow.shadow_type == MSO_SHADOW_TYPE.OUTER

    def it_provides_access_to_shadow_color_alpha(self):
        from pptx.oxml import parse_xml
        from pptx.oxml.ns import nsdecls

        # Shadow with 40% opacity (alpha = 0.4)
        spPr_xml = (
            '<p:spPr %s><a:effectLst>'
            '<a:outerShdw blurRad="50800">'
            '<a:srgbClr val="000000"><a:alpha val="40000"/></a:srgbClr>'
            '</a:outerShdw></a:effectLst></p:spPr>' % nsdecls("p", "a")
        )
        spPr = parse_xml(spPr_xml)
        shadow = ShadowFormat(spPr)
        assert shadow.color.alpha == pytest.approx(0.4)

    def it_can_set_shadow_color_alpha(self):
        shadow = ShadowFormat(element("p:spPr"))
        shadow.color.rgb = RGBColor(0, 0, 0)
        shadow.color.alpha = 0.5
        assert shadow.color.alpha == pytest.approx(0.5)

        # Setting to 1.0 removes the alpha element
        shadow.color.alpha = 1.0
        assert shadow.color.alpha == 1.0

    def it_reads_inner_shadow_color(self):
        from pptx.oxml import parse_xml
        from pptx.oxml.ns import nsdecls

        spPr_xml = (
            '<p:spPr %s><a:effectLst>'
            '<a:innerShdw blurRad="50800">'
            '<a:srgbClr val="00FF00"><a:alpha val="75000"/></a:srgbClr>'
            '</a:innerShdw></a:effectLst></p:spPr>' % nsdecls("p", "a")
        )
        spPr = parse_xml(spPr_xml)
        shadow = ShadowFormat(spPr)
        assert shadow.color.rgb == RGBColor(0x00, 0xFF, 0x00)
        assert shadow.color.alpha == pytest.approx(0.75)

    # fixtures -------------------------------------------------------

    @pytest.fixture(
        params=[
            ("p:spPr", True),
            ("p:spPr/a:effectLst", False),
            ("p:grpSpPr", True),
            ("p:grpSpPr/a:effectLst", False),
        ]
    )
    def inherit_get_fixture(self, request):
        cxml, expected_value = request.param
        shadow = ShadowFormat(element(cxml))
        return shadow, expected_value

    @pytest.fixture(
        params=[
            ("p:spPr{a:b=c}", False, "p:spPr{a:b=c}/a:effectLst"),
            ("p:grpSpPr{a:b=c}", False, "p:grpSpPr{a:b=c}/a:effectLst"),
            ("p:spPr{a:b=c}/a:effectLst", True, "p:spPr{a:b=c}"),
            ("p:grpSpPr{a:b=c}/a:effectLst", True, "p:grpSpPr{a:b=c}"),
            ("p:spPr", True, "p:spPr"),
            ("p:grpSpPr", True, "p:grpSpPr"),
            ("p:spPr/a:effectLst", False, "p:spPr/a:effectLst"),
            ("p:grpSpPr/a:effectLst", False, "p:grpSpPr/a:effectLst"),
        ]
    )
    def inherit_set_fixture(self, request):
        cxml, value, expected_cxml = request.param
        shadow = ShadowFormat(element(cxml))
        expected_value = xml(expected_cxml)
        return shadow, value, expected_value


class DescribeGlowFormat:
    @pytest.mark.parametrize(
        ("spPr_cxml", "expected_value"),
        [
            ("p:spPr", False),
            ("p:spPr/a:effectLst", False),
            ("p:spPr/a:effectLst/a:glow{rad=63500}", True),
        ],
    )
    def it_knows_whether_it_is_enabled(self, spPr_cxml: str, expected_value: bool):
        glow = GlowFormat(element(spPr_cxml))
        assert glow.enabled is expected_value

    @pytest.mark.parametrize(
        ("spPr_cxml", "expected_value"),
        [
            ("p:spPr", None),
            ("p:spPr/a:effectLst/a:glow{rad=63500}", 63500),
        ],
    )
    def it_knows_its_radius(self, spPr_cxml: str, expected_value):
        glow = GlowFormat(element(spPr_cxml))
        assert glow.radius == expected_value

    def it_adds_a_default_color_when_it_creates_a_glow(self):
        # -- `CT_GlowEffect` requires a color child; PowerPoint refuses a deck without one (#126)
        spPr = element("p:spPr{a:b=c}")
        glow = GlowFormat(spPr)
        glow.radius = 63500
        assert spPr.xml == xml(
            "p:spPr{a:b=c}/a:effectLst/a:glow{rad=63500}/a:schemeClr{val=accent1}/a:alpha{val=40000}"
        )

    def it_keeps_an_existing_color_when_it_changes_its_radius(self):
        spPr = element("p:spPr/a:effectLst/a:glow{rad=63500}/a:srgbClr{val=FFC000}")
        glow = GlowFormat(spPr)
        glow.radius = 127000
        assert spPr.xml == xml("p:spPr/a:effectLst/a:glow{rad=127000}/a:srgbClr{val=FFC000}")

    def it_repairs_a_colorless_glow_when_it_changes_its_radius(self):
        spPr = element("p:spPr/a:effectLst/a:glow{rad=63500}")
        glow = GlowFormat(spPr)
        glow.radius = 127000
        assert spPr.xml == xml(
            "p:spPr/a:effectLst/a:glow{rad=127000}/a:schemeClr{val=accent1}/a:alpha{val=40000}"
        )

    def it_provides_access_to_its_color(self):
        glow = GlowFormat(element("p:spPr/a:effectLst/a:glow{rad=63500}/a:srgbClr{val=FFC000}"))
        assert isinstance(glow.color, ColorFormat)
        assert glow.color.rgb == RGBColor(0xFF, 0xC0, 0x00)

    def it_can_change_its_color(self):
        spPr = element("p:spPr{a:b=c}")
        glow = GlowFormat(spPr)
        glow.color.rgb = RGBColor(0xFF, 0xC0, 0x00)
        assert glow.enabled is True
        assert spPr.xml == xml("p:spPr{a:b=c}/a:effectLst/a:glow/a:srgbClr{val=FFC000}")

    def it_round_trips_a_glow_through_a_saved_deck(self):
        # -- regression for #126: the deck from the issue's repro must carry a glow color --
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        box = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE, Emu(914400), Emu(1828800), Emu(2743200), Emu(1371600)
        )
        box.glow.radius = Pt(10)
        other = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.OVAL, Emu(914400), Emu(3657600), Emu(2743200), Emu(1371600)
        )
        other.glow.radius = Pt(5)
        other.glow.color.rgb = RGBColor(0xFF, 0xC0, 0x00)
        stream = io.BytesIO()
        prs.save(stream)

        stream.seek(0)
        shapes = Presentation(stream).slides[0].shapes
        reloaded_box, reloaded_other = shapes[-2], shapes[-1]
        assert reloaded_box.glow.radius == Pt(10)
        assert reloaded_box.glow.color.theme_color == MSO_THEME_COLOR.ACCENT_1
        assert reloaded_box.glow.color.alpha == pytest.approx(0.4)
        assert reloaded_other.glow.radius == Pt(5)
        assert reloaded_other.glow.color.rgb == RGBColor(0xFF, 0xC0, 0x00)

    def it_can_clear_itself(self):
        spPr = element("p:spPr/a:effectLst/a:glow{rad=63500}")
        glow = GlowFormat(spPr)
        glow.clear()
        assert glow.enabled is False


class DescribeReflectionFormat:
    @pytest.mark.parametrize(
        ("spPr_cxml", "expected_value"),
        [
            ("p:spPr", False),
            ("p:spPr/a:effectLst/a:reflection", True),
        ],
    )
    def it_knows_whether_it_is_enabled(self, spPr_cxml: str, expected_value: bool):
        reflection = ReflectionFormat(element(spPr_cxml))
        assert reflection.enabled is expected_value

    @pytest.mark.parametrize(
        ("spPr_cxml", "expected_value"),
        [
            ("p:spPr", None),
            ("p:spPr/a:effectLst/a:reflection{blurRad=6350}", 6350),
        ],
    )
    def it_knows_its_blur_radius(self, spPr_cxml: str, expected_value):
        reflection = ReflectionFormat(element(spPr_cxml))
        assert reflection.blur_radius == expected_value

    def it_can_change_its_blur_radius(self):
        spPr = element("p:spPr{a:b=c}")
        reflection = ReflectionFormat(spPr)
        reflection.blur_radius = 6350
        assert spPr.xml == xml("p:spPr{a:b=c}/a:effectLst/a:reflection{blurRad=6350}")

    def it_can_clear_itself(self):
        spPr = element("p:spPr/a:effectLst/a:reflection{blurRad=6350}")
        reflection = ReflectionFormat(spPr)
        reflection.clear()
        assert reflection.enabled is False


class DescribeSoftEdgeFormat:
    @pytest.mark.parametrize(
        ("spPr_cxml", "expected_value"),
        [
            ("p:spPr", False),
            ("p:spPr/a:effectLst/a:softEdge{rad=12700}", True),
        ],
    )
    def it_knows_whether_it_is_enabled(self, spPr_cxml: str, expected_value: bool):
        soft_edge = SoftEdgeFormat(element(spPr_cxml))
        assert soft_edge.enabled is expected_value

    @pytest.mark.parametrize(
        ("spPr_cxml", "expected_value"),
        [
            ("p:spPr", None),
            ("p:spPr/a:effectLst/a:softEdge{rad=12700}", 12700),
        ],
    )
    def it_knows_its_radius(self, spPr_cxml: str, expected_value):
        soft_edge = SoftEdgeFormat(element(spPr_cxml))
        assert soft_edge.radius == expected_value

    def it_can_change_its_radius(self):
        spPr = element("p:spPr{a:b=c}")
        soft_edge = SoftEdgeFormat(spPr)
        soft_edge.radius = 12700
        assert spPr.xml == xml("p:spPr{a:b=c}/a:effectLst/a:softEdge{rad=12700}")

    def it_can_clear_itself(self):
        spPr = element("p:spPr/a:effectLst/a:softEdge{rad=12700}")
        soft_edge = SoftEdgeFormat(spPr)
        soft_edge.clear()
        assert soft_edge.enabled is False
