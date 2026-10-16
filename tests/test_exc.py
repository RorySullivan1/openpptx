"""Unit-test suite for `pptx.exc`: the exception hierarchy callers rely on."""

from __future__ import annotations

import pytest

from pptx.exc import (
    ChartError,
    InvalidPackageError,
    InvalidTypeError,
    InvalidValueError,
    OpenPptxError,
    PackageError,
    ShapeError,
    SlideError,
    TableError,
)
from pptx.util import Inches


class DescribeExceptionHierarchy:
    """Each openpptx exception keeps the builtin base the API raised before it existed."""

    @pytest.mark.parametrize(
        ("exc_cls", "builtin"),
        [
            (InvalidValueError, ValueError),
            (InvalidTypeError, TypeError),
            (ShapeError, ValueError),
            (SlideError, ValueError),
            (ChartError, ValueError),
            (TableError, ValueError),
        ],
    )
    def it_is_both_an_OpenPptxError_and_the_builtin_it_replaced(self, exc_cls, builtin):
        assert issubclass(exc_cls, OpenPptxError)
        assert issubclass(exc_cls, builtin)

    def it_reports_an_unreadable_package_as_a_PackageError(self):
        assert issubclass(InvalidPackageError, PackageError)
        assert not issubclass(InvalidPackageError, ValueError)


class DescribeDeprecatedPyPptxError:
    """`PyPptxError`, the pre-openpptx name of the base class, still resolves, with a warning."""

    def it_is_the_same_class_as_OpenPptxError(self):
        with pytest.warns(DeprecationWarning, match="use OpenPptxError"):
            from pptx.exc import PyPptxError

        assert PyPptxError is OpenPptxError

    def it_warns_through_the_pptx_exceptions_alias_too(self):
        import pptx

        with pytest.warns(DeprecationWarning):
            assert pptx.exceptions.PyPptxError is OpenPptxError

    def but_it_still_raises_AttributeError_for_an_unknown_name(self):
        import pptx.exc

        with pytest.raises(AttributeError, match="no attribute 'Nope'"):
            pptx.exc.Nope  # noqa: B018


class DescribeValidationErrors:
    """Property setters reject bad values with the new classes, still catchable as builtins."""

    def it_raises_InvalidValueError_for_an_out_of_range_value(self):
        from pptx import Presentation

        prs = Presentation()
        textbox = prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(
            0, 0, Inches(1), Inches(1)
        )
        font = textbox.text_frame.paragraphs[0].font

        with pytest.raises(InvalidValueError):
            font.size = -1
        with pytest.raises(ValueError):
            font.size = -1

    def it_raises_InvalidTypeError_for_a_value_of_the_wrong_type(self):
        from pptx import Presentation

        prs = Presentation()
        textbox = prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(
            0, 0, Inches(1), Inches(1)
        )

        with pytest.raises(InvalidTypeError):
            textbox.left = "one inch"
        with pytest.raises(TypeError):
            textbox.left = "one inch"
