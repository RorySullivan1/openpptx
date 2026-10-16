"""Find/replace benchmarks (v0.7.0).

Times `Presentation.find_text()` and `Presentation.replace_text()` over the `large` synthetic deck:
100 slides x 25 text boxes x 5 paragraphs, each paragraph one run of text like
`slide=3 shape=12 p=4`, so 12,500 paragraphs in all.

- find, every paragraph matching (`shape=`), which reads every paragraph and builds a match each;
- replace with no match, the cost of the walk alone;
- replace in every paragraph, literal and as a regex with a template;
- `merge_runs()` over the whole deck once a word has been split into runs.

Each benchmark opens the deck fresh, so a replace never sees the result of an earlier one.
"""

from __future__ import annotations

import time
from pathlib import Path

import pytest

from pptx import Presentation
from pptx.text.search import iter_slide_text_frames


def _timed(fn) -> tuple[float, object]:
    t0 = time.perf_counter()
    result = fn()
    return time.perf_counter() - t0, result


@pytest.mark.perf
def bench_find_text_large(synth_large: Path) -> None:
    prs = Presentation(str(synth_large))
    elapsed, matches = _timed(lambda: prs.find_text("shape="))
    print(f"\n[bench_replace find_text large] {elapsed * 1000:.2f} ms  ({len(matches)} matches)")
    assert len(matches) == 12_500


@pytest.mark.perf
def bench_replace_text_no_match_large(synth_large: Path) -> None:
    prs = Presentation(str(synth_large))
    elapsed, count = _timed(lambda: prs.replace_text("no such text", "x"))
    print(f"\n[bench_replace replace_text no match large] {elapsed * 1000:.2f} ms")
    assert count == 0


@pytest.mark.perf
def bench_replace_text_every_paragraph_large(synth_large: Path) -> None:
    prs = Presentation(str(synth_large))
    elapsed, count = _timed(lambda: prs.replace_text("shape=", "shp:"))
    print(f"\n[bench_replace replace_text literal large] {elapsed * 1000:.2f} ms  ({count})")
    assert count == 12_500


@pytest.mark.perf
def bench_replace_text_regex_large(synth_large: Path) -> None:
    prs = Presentation(str(synth_large))
    elapsed, count = _timed(lambda: prs.replace_text(r"p=(\d+)", r"para \1", regex=True))
    print(f"\n[bench_replace replace_text regex large] {elapsed * 1000:.2f} ms  ({count})")
    assert count == 12_500


@pytest.mark.perf
def bench_merge_runs_large(synth_large: Path) -> None:
    prs = Presentation(str(synth_large))
    # -- split each paragraph's run in two, with identical formatting, so each can merge back --
    for slide in prs.slides:
        for _, text_frame in iter_slide_text_frames(slide):
            for paragraph in text_frame.paragraphs:
                (run,) = paragraph.runs
                text = run.text
                run.text = text[:5]
                paragraph.add_run().text = text[5:]

    def merge() -> int:
        return sum(
            text_frame.merge_runs()
            for slide in prs.slides
            for _, text_frame in iter_slide_text_frames(slide)
        )

    elapsed, merged = _timed(merge)
    print(f"\n[bench_replace merge_runs large] {elapsed * 1000:.2f} ms  ({merged} merged)")
    assert merged == 12_500
