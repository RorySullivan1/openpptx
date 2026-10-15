# pyright: reportPrivateUsage=false

"""Unit-test suite for `pptx.parts.chartex` module."""

from __future__ import annotations

import pytest

from pptx import Presentation
from pptx.chart.chartex import ChartEx
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.opc.constants import CONTENT_TYPE as CT
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part, PartFactory
from pptx.opc.packuri import PackURI
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.parts.chartex import ChartExPart
from pptx.parts.embeddedpackage import EmbeddedXlsxPart

from ..unitutil.chartex import WATERFALL_XML, xlsx_column
from ..unitutil.mock import function_mock, instance_mock

_CS_NS = "http://schemas.microsoft.com/office/drawing/2012/chartStyle"


class DescribeChartExPart:
    """Unit-test suite for `pptx.parts.chartex.ChartExPart` objects."""

    def it_is_registered_for_the_chartex_content_type(self):
        assert PartFactory.part_type_for[CT.OFC_CHART_EX] is ChartExPart

    def it_can_create_a_new_chart_with_its_workbook(self):
        package = Presentation().part.package
        chart_data = CategoryChartData()
        chart_data.categories = ["Start", "End"]
        chart_data.add_series("Flow", (10, 12))

        part = ChartExPart.new(XL_CHART_TYPE.WATERFALL, chart_data, package)

        assert part.partname == PackURI("/ppt/charts/chartEx1.xml")
        assert part.content_type == CT.OFC_CHART_EX
        xlsx_part = part.xlsx_part
        assert isinstance(xlsx_part, EmbeddedXlsxPart)
        assert xlsx_column(xlsx_part.blob, "B") == [10.0, 12.0]
        externalData = part.chartSpace.xpath("./cx:chartData/cx:externalData")[0]
        assert externalData.get(qn("cx:autoUpdate")) == "0"
        assert part.chartex.chart_type is XL_CHART_TYPE.WATERFALL

    @pytest.mark.parametrize(
        ("chart_type", "style_id"),
        [
            (XL_CHART_TYPE.WATERFALL, "395"),
            (XL_CHART_TYPE.HISTOGRAM, "366"),
            (XL_CHART_TYPE.BOX_WHISKER, "406"),
            (XL_CHART_TYPE.TREEMAP, "410"),
            (XL_CHART_TYPE.SUNBURST, "381"),
            (XL_CHART_TYPE.FUNNEL, "419"),
        ],
    )
    def it_adds_the_style_parts_powerpoint_requires(self, chart_type, style_id: str):
        # -- PowerPoint refuses a chartex chart lacking either part (#128); the ids are the ones
        # -- PowerPoint itself writes for a new chart of each type
        package = Presentation().part.package
        chart_data = CategoryChartData()
        chart_data.categories = ["A", "B"]
        chart_data.add_series("S1", (10, 12))

        part = ChartExPart.new(chart_type, chart_data, package)

        style_part = part.part_related_by(RT.CHART_STYLE)
        assert style_part.partname == PackURI("/ppt/charts/style1.xml")
        assert style_part.content_type == CT.OFC_CHART_STYLE
        style = parse_xml(style_part.blob)
        assert style.tag == "{%s}chartStyle" % _CS_NS
        assert style.get("id") == style_id
        colors_part = part.part_related_by(RT.CHART_COLOR_STYLE)
        assert colors_part.partname == PackURI("/ppt/charts/colors1.xml")
        assert colors_part.content_type == CT.OFC_CHART_COLORS
        colors = parse_xml(colors_part.blob)
        assert colors.tag == "{%s}colorStyle" % _CS_NS
        assert colors.get("id") == "10"

    @pytest.mark.parametrize(
        ("chart_type", "style_id"),
        [(XL_CHART_TYPE.WATERFALL, "395"), (XL_CHART_TYPE.FUNNEL, "419")],
    )
    def it_adds_the_style_parts_a_loaded_chart_is_missing(self, chart_type, style_id: str):
        # -- a chartex chart written by pypptx 0.5.0 has neither part; PowerPoint refuses it (#132)
        package = Presentation().part.package
        blob = WATERFALL_XML.replace('layoutId="waterfall"', 'layoutId="%s"' % (
            "waterfall" if chart_type is XL_CHART_TYPE.WATERFALL else "funnel"
        ))
        part = ChartExPart.load(
            PackURI("/ppt/charts/chartEx1.xml"), CT.OFC_CHART_EX, package, blob.encode("utf-8")
        )

        part.ensure_style_parts()

        style_part = part.part_related_by(RT.CHART_STYLE)
        assert style_part.content_type == CT.OFC_CHART_STYLE
        assert parse_xml(style_part.blob).get("id") == style_id
        colors_part = part.part_related_by(RT.CHART_COLOR_STYLE)
        assert colors_part.content_type == CT.OFC_CHART_COLORS
        assert parse_xml(colors_part.blob).get("id") == "10"

    def it_adds_style_parts_even_when_its_chart_type_is_unknown(self):
        # -- a chart pypptx can't classify must still open in PowerPoint --
        package = Presentation().part.package
        blob = (
            b'<cx:chartSpace xmlns:cx="http://schemas.microsoft.com/office/drawing/2014/chartex"'
            b"/>"
        )
        part = ChartExPart.load(
            PackURI("/ppt/charts/chartEx1.xml"), CT.OFC_CHART_EX, package, blob
        )

        part.ensure_style_parts()

        style = parse_xml(part.part_related_by(RT.CHART_STYLE).blob)
        assert style.tag == "{%s}chartStyle" % _CS_NS
        colors = parse_xml(part.part_related_by(RT.CHART_COLOR_STYLE).blob)
        assert colors.tag == "{%s}colorStyle" % _CS_NS

    def it_leaves_style_parts_that_are_already_there_alone(self):
        package = Presentation().part.package
        chart_data = CategoryChartData()
        chart_data.categories = ["A", "B"]
        chart_data.add_series("S1", (10, 12))
        part = ChartExPart.new(XL_CHART_TYPE.WATERFALL, chart_data, package)
        rels_before = dict(part.rels)

        part.ensure_style_parts()

        assert dict(part.rels) == rels_before

    def it_does_not_parse_a_chart_whose_style_parts_are_present(self, request):
        parse_xml_ = function_mock(request, "pptx.opc.package.parse_xml")
        package_ = instance_mock(request, "pptx.package.Package")
        part = ChartExPart.load(
            PackURI("/ppt/charts/chartEx1.xml"),
            CT.OFC_CHART_EX,
            package_,
            WATERFALL_XML.encode("utf-8"),
        )
        style_part_ = instance_mock(request, Part)
        colors_part_ = instance_mock(request, Part)
        part.relate_to(style_part_, RT.CHART_STYLE)
        part.relate_to(colors_part_, RT.CHART_COLOR_STYLE)

        part.ensure_style_parts()

        parse_xml_.assert_not_called()

    def it_provides_access_to_the_chart(self, request):
        package_ = instance_mock(request, "pptx.package.Package")
        part = ChartExPart.load(
            PackURI("/ppt/charts/chartEx1.xml"),
            CT.OFC_CHART_EX,
            package_,
            WATERFALL_XML.encode("utf-8"),
        )

        chartex = part.chartex

        assert isinstance(chartex, ChartEx)
        assert chartex.part is part
        assert chartex.chart_type is XL_CHART_TYPE.WATERFALL

    def it_has_no_workbook_when_there_is_no_external_data_reference(self, request):
        package_ = instance_mock(request, "pptx.package.Package")
        blob = (
            b'<cx:chartSpace xmlns:cx="http://schemas.microsoft.com/office/drawing/2014/chartex"'
            b"/>"
        )
        part = ChartExPart.load(
            PackURI("/ppt/charts/chartEx1.xml"), CT.OFC_CHART_EX, package_, blob
        )

        assert part.xlsx_part is None

    def it_never_parses_when_only_the_blob_is_read_back(self, request):
        parse_xml_ = function_mock(request, "pptx.opc.package.parse_xml")
        package_ = instance_mock(request, "pptx.package.Package")
        blob = WATERFALL_XML.encode("utf-8")

        part = ChartExPart.load(
            PackURI("/ppt/charts/chartEx1.xml"), CT.OFC_CHART_EX, package_, blob
        )

        assert part.blob == blob
        parse_xml_.assert_not_called()
